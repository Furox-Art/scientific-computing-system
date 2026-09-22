#!/usr/bin/env python3
import hashlib, json, os, tempfile, urllib.request, zipfile
from pathlib import Path

URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-021-23540-y/MediaObjects/41467_2021_23540_MOESM4_ESM.zip"
OUT=Path("pereira_audit/NATURE_SOURCE_DATA_AUDIT_2026-09-22.json")

def sha256(p):
    h=hashlib.sha256()
    with open(p,"rb") as f:
        for b in iter(lambda:f.read(4*1024*1024),b""): h.update(b)
    return h.hexdigest()

def xls_preview(raw):
    import io, xlrd
    book=xlrd.open_workbook(file_contents=raw)
    sheets=[]
    for ws in book.sheets():
        rows=[]
        for r in range(min(ws.nrows,300)):
            rows.append([ws.cell_value(r,col) for col in range(ws.ncols)])
        sheets.append({"title":ws.name,"max_row":ws.nrows,"max_column":ws.ncols,"rows":rows})
    return sheets

def xlsx_preview(raw):
    import io
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        rows=[]
        for i,row in enumerate(ws.iter_rows(values_only=True),1):
            rows.append(list(row))
            if i>=200: break
        sheets.append({"title":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,"rows":rows})
    return sheets

with tempfile.TemporaryDirectory(prefix="pereira-nature-") as td:
    zpath=Path(td)/"source.zip"
    req=urllib.request.Request(URL,headers={"User-Agent":"Mozilla/5.0"})
    with urllib.request.urlopen(req,timeout=300) as r, open(zpath,"wb") as f:
        while True:
            b=r.read(4*1024*1024)
            if not b: break
            f.write(b)
    audit={
      "status":"PEREIRA2021_NATURE_SOURCE_DATA_AUDITED",
      "source_url":URL,
      "zip_bytes":zpath.stat().st_size,
      "zip_sha256":sha256(zpath),
      "members":[],
      "tabular_previews":[]
    }
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist():
            audit["members"].append({"name":info.filename,"bytes":info.file_size,"compressed_bytes":info.compress_size})
            low=info.filename.lower()
            if low.endswith(".xls") and not info.filename.startswith("__MACOSX/"):
                raw=z.read(info)
                try:
                    sheets=xls_preview(raw)
                    audit["tabular_previews"].append({"member":info.filename,"type":"xls","sheets":sheets})
                except Exception as e:
                    audit["tabular_previews"].append({"member":info.filename,"type":"xls","error":repr(e)})
            elif low.endswith(".xlsx"):
                raw=z.read(info)
                try:
                    sheets=xlsx_preview(raw)
                    audit["tabular_previews"].append({"member":info.filename,"type":"xlsx","sheets":sheets})
                except Exception as e:
                    audit["tabular_previews"].append({"member":info.filename,"type":"xlsx","error":repr(e)})
            elif low.endswith(".m") and ("fig4" in low or "config_model" in low):
                raw=z.read(info).decode("utf-8","replace")
                audit["tabular_previews"].append({"member":info.filename,"type":"matlab","lines":raw.splitlines()[:1000]})
            elif low.endswith((".csv",".tsv",".txt")) and info.file_size<5_000_000:
                raw=z.read(info).decode("utf-8","replace")
                audit["tabular_previews"].append({"member":info.filename,"type":"text","lines":raw.splitlines()[:500]})
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(audit,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({
      "zip_bytes":audit["zip_bytes"],
      "zip_sha256":audit["zip_sha256"],
      "members":audit["members"],
      "tabular_members":[x["member"] for x in audit["tabular_previews"]],
      "fig4h":[x for x in audit["tabular_previews"] if "fig.4h" in x["member"].lower()],
      "fig4_matlab":[x for x in audit["tabular_previews"] if x["member"].lower().endswith("fig4.m")]
    },indent=2))
