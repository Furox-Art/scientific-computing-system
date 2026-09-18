#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, re, sys, traceback
from pathlib import Path
import requests
from docx import Document

URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41598-025-12695-z/MediaObjects/41598_2025_12695_MOESM1_ESM.docx"
OUTDIR=Path(".public-runner/p07")
DOCX=OUTDIR/"41598_2025_12695_MOESM1_ESM.docx"
RESULT=OUTDIR/"supplement_mapping_scan.json"

SUBJ_RE=re.compile(r"(?i)\b(?:sub[-_ ]?|SD[-_ ]?)?10\d{2}\b")
KEY_RE=re.compile(r"(?i)\b(no experience|no information|experience|awakening|confidence score|dream report|report category)\b")

def clean(s):
    return " ".join((s or "").split())

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    result={"status":"STARTED","url":URL,"mapping_rows":[],"relevant_paragraphs":[],"tables":[]}
    try:
        r=requests.get(URL,timeout=(30,180),allow_redirects=True,headers={"User-Agent":"github-actions-public-validation/1.0"})
        result["http_status"]=r.status_code
        result["final_url"]=r.url
        if r.status_code!=200: raise RuntimeError(f"HTTP {r.status_code}")
        DOCX.write_bytes(r.content)
        result["bytes"]=len(r.content)
        result["sha256"]=hashlib.sha256(r.content).hexdigest()
        doc=Document(DOCX)

        for i,p in enumerate(doc.paragraphs):
            t=clean(p.text)
            if not t: continue
            if KEY_RE.search(t) or (SUBJ_RE.search(t) and len(t)<=500):
                result["relevant_paragraphs"].append({"paragraph_index":i,"text":t[:500]})

        for ti,table in enumerate(doc.tables):
            rows=[[clean(c.text) for c in row.cells] for row in table.rows]
            header=rows[0] if rows else []
            tab={"table_index":ti,"rows":len(rows),"cols":max([len(x) for x in rows],default=0),"header":header}
            hit_count=0
            for ri,row in enumerate(rows):
                joined=" | ".join(row)
                if SUBJ_RE.search(joined) or (KEY_RE.search(joined) and ri<5):
                    rec={"table_index":ti,"row_index":ri,"cells":row}
                    result["mapping_rows"].append(rec)
                    hit_count+=1
            tab["candidate_rows"]=hit_count
            result["tables"].append(tab)

        result["mapping_row_count"]=len(result["mapping_rows"])
        result["relevant_paragraph_count"]=len(result["relevant_paragraphs"])
        result["status"]="SUPPLEMENT_DOWNLOADED_AND_MAPPING_CANDIDATES_EXTRACTED"
        RESULT.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
        print(json.dumps({k:result[k] for k in ["status","bytes","sha256","mapping_row_count","relevant_paragraph_count","tables"]},indent=2))
        return 0
    except Exception as e:
        result["status"]="SUPPLEMENT_MAPPING_SCAN_FAILED"
        result["error"]={"type":type(e).__name__,"message":str(e),"traceback":traceback.format_exc()}
        RESULT.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
        print(json.dumps(result,indent=2)); return 2

if __name__=="__main__":
    sys.exit(main())
